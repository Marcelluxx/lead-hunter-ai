"""Authorized reference-only XLSX generation and delivery; no provider I/O."""
from __future__ import annotations

from collections.abc import Callable, Iterable
from datetime import datetime, timezone
from io import BytesIO
import os
from pathlib import Path
import tempfile
from urllib.parse import quote, urlencode

from openpyxl import Workbook
from openpyxl.packaging.core import DocumentProperties

from src.application.feature_access import FeatureAccessService
from src.application.rating_filters import RatingFilterGuard
from src.domain.feature_licenses import FeatureAction, FeatureContext
from src.domain.place_references import GooglePlaceReference, ReferenceExportError, normalize_references


def _serialize_reference_workbook(references: tuple[GooglePlaceReference, ...]) -> bytes:
    workbook = Workbook()
    # Fixed neutral metadata: no query, machine identity or provider content.
    timestamp = datetime(2000, 1, 1, tzinfo=timezone.utc)
    workbook.properties = DocumentProperties(creator='Lead Hunter', lastModifiedBy='Lead Hunter',
                                              created=timestamp, modified=timestamp)
    sheet = workbook.active
    sheet.title = 'Riferimenti'
    sheet.append(['Place ID', 'Link Google Maps'])
    for reference in references:
        query = urlencode({'api': '1', 'query': 'attività', 'query_place_id': reference.place_id},
                          quote_via=quote)
        link = 'https://www.google.com/maps/search/?' + query
        sheet.append([reference.place_id, link])
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'
        sheet.cell(sheet.max_row, 2).hyperlink = link
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:B{sheet.max_row}'
    sheet.column_dimensions['A'].width = 45
    sheet.column_dimensions['B'].width = 90
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class ReferenceExportService:
    def __init__(self, access: FeatureAccessService, context_factory: Callable[[], FeatureContext]):
        self.access = access
        self.context_factory = context_factory

    def require_access(self, *, action: FeatureAction = FeatureAction.EXECUTE) -> None:
        self.access.require(self.context_factory(), 'export.no_website', action=action)

    def export_bytes(self, references: Iterable[GooglePlaceReference], *,
                     origin: RatingFilterGuard | None = None) -> bytes:
        if origin is not None:
            origin.require_view()
        self.require_access()
        normalized = normalize_references(references)
        data = _serialize_reference_workbook(normalized)
        self.require_access(action=FeatureAction.VIEW)
        if origin is not None:
            origin.require_view()
        return data

    def save(self, references: Iterable[GooglePlaceReference], destination: Path, *,
             origin: RatingFilterGuard | None = None) -> int:
        if origin is not None:
            origin.require_view()
        self.require_access()
        normalized = normalize_references(references)
        data = _serialize_reference_workbook(normalized)
        destination = Path(destination)
        temporary: Path | None = None
        try:
            try:
                destination.parent.mkdir(parents=True, exist_ok=True)
                with tempfile.NamedTemporaryFile(mode='wb', dir=destination.parent,
                                                  prefix='.reference-export-', delete=False) as handle:
                    temporary = Path(handle.name)
                    handle.write(data)
                    handle.flush()
                    os.fsync(handle.fileno())
                self.require_access(action=FeatureAction.VIEW)
                if origin is not None:
                    origin.require_view()
                os.replace(temporary, destination)
                return len(normalized)
            finally:
                if temporary is not None:
                    temporary.unlink(missing_ok=True)
        except OSError:
            raise ReferenceExportError('reference_export_write_failed') from None
