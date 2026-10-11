"""Actual HTTP -> least-privilege admission -> worker -> protected delivery on PostgreSQL."""
import unittest
from dataclasses import replace
from io import BytesIO
from uuid import UUID
from unittest.mock import patch

from openpyxl import load_workbook

from src.application.job_execution import JobExecutionService
from src.infrastructure.crypto import AccessTokenService, PasswordService, SecretCipher
from src.application.authentication import AuthenticationService
from tests.integration.postgres_helpers import POSTGRES_AVAILABLE, databases, delete_fixture
from tests import test_commercial_job_api as api_tests


@unittest.skipUnless(POSTGRES_AVAILABLE, 'PostgreSQL integration URLs not configured')
class CommercialDeliveryPostgresTests(unittest.TestCase):
    def test_http_worker_and_reopened_export_with_actual_app_worker_roles(self):
        owner, app, worker = databases()
        private, public = AccessTokenService.generate_key_pair()
        cipher = SecretCipher(SecretCipher.generate_key())
        authentication = AuthenticationService(passwords=PasswordService(), cipher=cipher,
            tokens=AccessTokenService(private_key_pem=private, public_key_pem=public, issuer='fixture', audience='fixture'))
        fixture = api_tests.CommercialJobApiTests()
        try:
            with patch('tests.test_job_execution.platform_fixture', return_value=(owner, PasswordService(), cipher, authentication)):
                fixture.setUp()
            fx = fixture.fx
            app_execution = JobExecutionService(database=app, installation_id=fx.installation, access_factory=fx.access)
            worker_execution = JobExecutionService(database=worker, installation_id=fx.installation, access_factory=fx.access)
            fixture.client.app.state.runtime = replace(fixture.app_runtime, database=app, job_execution=app_execution)
            fixture.worker.runtime.execution = worker_execution
            fx.activate(('export.no_website',))
            admitted = fixture.create()
            self.assertEqual(admitted.status_code, 202, admitted.text)
            job_id = UUID(admitted.json()['id'])
            fixture.worker.runtime.execute(job_id)
            url = fixture.url + '/' + str(job_id)
            results = fixture.client.get(url + '/results', headers=fixture.headers)
            self.assertEqual(results.status_code, 200, results.text)
            self.assertEqual(results.json()['references'][0]['place_id'], 'place-1')
            exported = fixture.client.get(url + '/export', headers=fixture.headers)
            self.assertEqual(exported.status_code, 200)
            self.assertEqual(load_workbook(BytesIO(exported.content)).active.cell(2, 1).value, 'place-1')
            fx.clock.set(100)
            self.assertEqual(fixture.client.get(url + '/export', headers=fixture.headers).status_code, 403)
        finally:
            if hasattr(fixture, 'fx'):
                delete_fixture(owner, fixture.fx.workspace.id, fixture.fx.user.id)
            fixture.doCleanups()
            for database in (owner, app, worker): database.engine.dispose()
