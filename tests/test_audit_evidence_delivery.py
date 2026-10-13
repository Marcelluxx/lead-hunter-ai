import asyncio
import copy
import json
import unittest
from datetime import datetime
from hashlib import sha256
from io import BytesIO
from uuid import UUID

import httpx
from openai import OpenAI
from openpyxl import load_workbook
from sqlalchemy import select

from src.application.provenance import build_verified_lead
from src.domain.audit import AuditValidationError
from src.domain.audit_evidence import AuditEvidenceSet
from src.domain.crawl import PageEvidence
from src.exporter import DataExporter
from src.infrastructure.models import LeadAttributeModel, EvidenceModel
from tests.test_compliant_pipeline import _Crawler
from tests.test_audit_evidence import QUOTE, finding
from tests.test_audit_evidence_runtime import LEGACY


def crawl_fixture(base=None):
    crawl = base if base is not None else asyncio.run(_Crawler().crawl('https://official.example.test/'))
    crawl.pages[crawl.url] += QUOTE
    crawl.pages['https://official.example.test/about'] = 'Informazioni sullo studio. ' * 20 + QUOTE
    crawl.evidence = [PageEvidence.from_content(requested_url=url, final_url=url, status_code=200,
        content_type='text/html', content=text) for url, text in crawl.pages.items()]
    return crawl


def report_fixture(crawl):
    context = AuditEvidenceSet.from_pages(crawl.pages)
    catalog = json.loads(context.to_prompt_payload())
    return context.validate_findings({'findings': [finding(catalog['pages'][1]['evidence_id'])]})


class AuditEvidenceDeliveryTests(unittest.TestCase):
    def test_source_path_identifiers_are_redacted_in_public_evidence(self):
        for path in ['/mario.rossi@site.test', '/mario.rossi%40site.test', '/api_key=PRIVATE_PATH_KEY']:
            report = AuditEvidenceSet.from_pages({'https://site.test' + path: 'Contenuto ufficiale ' * 30}).empty_report()
            public = json.dumps(report)
            self.assertNotIn('mario.rossi', public)
            self.assertNotIn('PRIVATE_PATH_KEY', public)
            self.assertIn('REDACTED_PATH', public)

    def test_structurally_valid_oversized_report_is_blocked_before_export(self):
        crawl = crawl_fixture(); report = report_fixture(crawl)
        original = report['findings'][0]
        report['findings'] = []
        for i in range(20):
            item = copy.deepcopy(original); item['id'] = f'f{i + 1}'; item['description'] = 'È' * 990 + str(i)
            report['findings'].append(item)
        self.assertGreater(len(json.dumps(report, ensure_ascii=False).encode()), 24000)
        with self.assertRaises(AuditValidationError):
            build_verified_lead(crawl, search_keyword='Studio', audit={**LEGACY, 'audit_report': report})

    def test_builder_binds_original_hash_and_actual_page_time_and_preserves_local_xlsx(self):
        crawl = crawl_fixture()
        lead = build_verified_lead(crawl, search_keyword='Studio', audit={**LEGACY, 'audit_report': report_fixture(crawl)})
        report = lead.audit_report
        citation = report['findings'][0]['citations'][0]
        self.assertEqual(citation['retrieved_at'], datetime.fromisoformat(crawl.evidence[1].retrieved_at).isoformat())
        self.assertEqual(citation['content_sha256'], sha256(crawl.pages[citation['source_url']].encode()).hexdigest())
        book = load_workbook(BytesIO(DataExporter.export_bytes([lead], suppression_checker=lambda _: False)))
        row = dict(zip([x.value for x in book.active[1]], [x.value for x in book.active[2]]))
        self.assertEqual(json.loads(row['Audit Report']), report)
        self.assertEqual(lead.provenance['audit_report'].data_classification, 'website_audit')
        self.assertNotIn('full_prompt', row['Audit Report'])

    def test_builder_refuses_mutated_report_and_stale_crawl_evidence(self):
        crawl = crawl_fixture(); original = report_fixture(crawl)
        for key, value in [('excerpt', 'This quotation was never on the page.'), ('quote_sha256', 'a' * 64)]:
            bad = copy.deepcopy(original); bad['findings'][0]['citations'][0][key] = value
            with self.subTest(key=key), self.assertRaises(AuditValidationError):
                build_verified_lead(crawl, search_keyword='Studio', audit={**LEGACY, 'audit_report': bad})
        crawl.pages[crawl.url] += 'Mutated after evidence collection.'
        with self.assertRaises(AuditValidationError):
            build_verified_lead(crawl, search_keyword='Studio', audit={**LEGACY, 'audit_report': original})

    def test_legacy_audit_is_declared_uncited_without_inventing_findings(self):
        lead = build_verified_lead(crawl_fixture(), search_keyword='Studio', audit=LEGACY)
        self.assertEqual(lead.audit_report['status'], 'legacy_uncited')
        self.assertEqual(lead.audit_report['findings'], [])

    def test_failed_local_audit_has_no_score_or_bad_quality_highlight(self):
        crawl = crawl_fixture()
        report = AuditEvidenceSet.from_pages(crawl.pages).empty_report('failed')
        lead = build_verified_lead(crawl, search_keyword='Studio', audit={**LEGACY, 'website_score': None, 'audit_report': report})
        book = load_workbook(BytesIO(DataExporter.export_bytes([lead], suppression_checker=lambda _: False)))
        headers = [cell.value for cell in book.active[1]]
        score = book.active.cell(2, headers.index('Website Score') + 1)
        self.assertIsNone(score.value)
        self.assertNotEqual(score.fill.fgColor.rgb, '00FEE2E2')


class AuditEvidenceServerDeliveryTests(unittest.TestCase):
    def setUp(self):
        from tests.test_commercial_job_api import CommercialJobApiTests
        self.fx = CommercialJobApiTests(); self.fx.setUp(); self.addCleanup(self.fx.doCleanups)

    def cited_job(self):
        class Crawler(_Crawler):
            async def crawl(self, url): return crawl_fixture(await super().crawl(url))
        self.fx.worker.runtime.crawler_factory = lambda **_: Crawler()
        def transport(request):
            body = json.loads(request.content)
            if body['model'] == 'clean': content = 'Cleaning paraphrase ' * 30
            else:
                catalog = json.loads(body['messages'][1]['content'].split('ORIGINAL_AUDIT_EVIDENCE:\n')[1])
                content = json.dumps({**LEGACY, 'findings': [finding(catalog['pages'][1]['evidence_id'])]})
            return httpx.Response(200, json={'id': 'fixture', 'object': 'chat.completion', 'created': 1,
                'model': body['model'], 'choices': [{'index': 0, 'message': {'role': 'assistant', 'content': content}, 'finish_reason': 'stop'}]})
        client = httpx.Client(transport=httpx.MockTransport(transport)); self.addCleanup(client.close)
        sdk = OpenAI(api_key='fixture', base_url='https://llm.test/', http_client=client, max_retries=0); self.addCleanup(sdk.close)
        self.fx.worker.runtime.auditor_client = sdk
        response = self.fx.create(kind='audit.website', parameters={'urls': ['https://official.example.test/'], 'max_pages': 2})
        self.assertEqual(response.status_code, 202, response.text)
        job_id = UUID(response.json()['id']); self.fx.worker.runtime.execute(job_id)
        url = self.fx.url + '/' + str(job_id)
        return job_id, url

    def test_sdk_worker_sql_http_and_workbook_share_bounded_public_report(self):
        job_id, url = self.cited_job()
        response = self.fx.client.get(url + '/results', headers=self.fx.headers)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['state'], 'completed')
        lead = response.json()['leads'][0]; report = lead['audit_report']
        self.assertEqual(report['findings'][0]['citations'][0]['excerpt'], QUOTE)
        self.assertEqual(report['findings'][0]['status'], 'inconclusive')
        self.assertTrue(all(any(e['source_url'] == x['source_url'] and e['content_sha256'] == x['content_sha256']
            for e in lead['evidence']) for x in report['evidence']))
        with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
            attribute = session.scalar(select(LeadAttributeModel).where(LeadAttributeModel.field_name == 'audit_report'))
            self.assertEqual(json.loads(attribute.value_text), report)
        export = self.fx.client.get(url + '/export', headers=self.fx.headers)
        self.assertEqual(export.status_code, 200, export.text)
        book = load_workbook(BytesIO(export.content))
        row = dict(zip([x.value for x in book.active[1]], [x.value for x in book.active[2]]))
        self.assertEqual(json.loads(row['audit_report']), report)
        self.assertLessEqual(len(row['audit_report'].encode()), 24000)
        self.assertNotIn('PRIVATE GOOGLE', repr(row))

    def test_corrupted_report_json_or_unbound_evidence_fails_safely(self):
        _, url = self.cited_job()
        with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
            attribute = session.scalar(select(LeadAttributeModel).where(LeadAttributeModel.field_name == 'audit_report'))
            self.assertIsNotNone(attribute)
            original = attribute.value_text; attribute.value_text = '{PRIVATE_SQL_SENTINEL'
        response = self.fx.client.get(url + '/results', headers=self.fx.headers)
        self.assertEqual(response.json(), {'detail': 'job_audit_report_invalid'})
        self.assertNotIn('PRIVATE_SQL_SENTINEL', response.text)
        with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
            attribute = session.scalar(select(LeadAttributeModel).where(LeadAttributeModel.field_name == 'audit_report'))
            attribute.value_text = original
            evidence = session.scalar(select(EvidenceModel).where(EvidenceModel.source_url.endswith('/about')))
            self.assertIsNotNone(evidence); evidence.content_sha256 = 'a' * 64
        response = self.fx.client.get(url + '/export', headers=self.fx.headers)
        self.assertEqual(response.json(), {'detail': 'job_audit_report_invalid'})

    def test_historical_record_without_report_is_uncited_and_legacy_attributes_stay_strings(self):
        _, url = self.cited_job()
        with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
            attribute = session.scalar(select(LeadAttributeModel).where(LeadAttributeModel.field_name == 'audit_report'))
            self.assertIsNotNone(attribute); session.delete(attribute)
        response = self.fx.client.get(url + '/results', headers=self.fx.headers)
        lead = response.json()['leads'][0]
        self.assertEqual(lead['audit_report']['status'], 'legacy_uncited')
        self.assertEqual(lead['audit_report']['findings'], [])
        self.assertIsInstance(lead['attributes']['website_score'], str)
