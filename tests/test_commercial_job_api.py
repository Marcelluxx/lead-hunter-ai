import json
import unittest
from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4
from unittest.mock import patch
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from src.application.authentication import AuthenticationService
from src.application.budgets import BudgetService
from src.application.jobs import JobService
from src.infrastructure.crypto import AccessTokenService, PasswordService, SecretCipher
from src.infrastructure.models import JobModel, UsageBudgetModel
from src.web.app import create_app
from src.web.dependencies import WebRuntime
from tests import test_commercial_worker as worker_tests
from tests.test_api_security import Publisher


class CommercialJobApiTests(unittest.TestCase):
    def setUp(self):
        self.worker = worker_tests.CommercialWorkerTests(); self.worker.setUp(); self.addCleanup(self.worker.doCleanups)
        self.fx = self.worker.fx
        private, public = AccessTokenService.generate_key_pair()
        authentication = AuthenticationService(passwords=PasswordService(), cipher=SecretCipher(SecretCipher.generate_key()),
            tokens=AccessTokenService(private_key_pem=private, public_key_pem=public, issuer='test', audience='test'))
        self.publisher = Publisher()
        self.app_runtime = WebRuntime(database=self.fx.database, authentication=authentication,
            jobs=JobService(budgets=BudgetService(), publisher=self.publisher), budgets=BudgetService(),
            credentials=self.worker.credentials, suppression=self.worker.runtime.contacts.suppression,
            job_execution=self.fx.service, job_prices=self.fx.prices, diagnostic_root=Path(self.worker.temp.name))
        self.client = TestClient(create_app(self.app_runtime), base_url='https://testserver')
        self.addCleanup(self.client.close)
        login = self.client.post('/auth/login', json={'email': 'operator@example.test', 'password': 'correct horse battery staple'})
        token = login.json()['access_token']
        headers = {'Authorization': 'Bearer ' + token}
        code = self.client.post('/auth/mfa/enrollment', headers=headers).json()['recovery_codes'][0]
        token = self.client.post('/auth/mfa/confirm', headers=headers, json={'code': code}).json()['access_token']
        self.headers = {'Authorization': 'Bearer ' + token}
        with self.fx.database.session() as session:
            self.fx.auth = authentication.authenticate_access_token(session, token)
        self.url = f'/workspaces/{self.fx.workspace.id}/jobs'

    def create(self, *, kind='discover.no_website', parameters=None, key=None):
        return self.client.post(self.url, headers={**self.headers, 'Idempotency-Key': key or str(uuid4())},
            json={'kind': kind, 'parameters': parameters or {'lat': 0, 'lng': 0, 'keywords': ['dentista'], 'max_leads': 1}, 'estimated_cost': '0'})

    def _named_contact_job(self):
        from datetime import datetime, timezone
        from src.application.privacy_policy import WorkspacePrivacyPolicyService
        from src.domain.privacy import WorkspacePrivacyPolicy
        from src.domain.contacts import ContactPoint, ContactExtractionMethod
        from src.infrastructure.models import LeadModel
        response = self.create(kind='audit.website', parameters={'urls': ['https://official.example.test/'], 'max_pages': 1})
        self.assertEqual(response.status_code, 202, response.text)
        job_id = UUID(response.json()['id']); self.worker.runtime.execute(job_id)
        with self.fx.database.session(self.fx.workspace.id) as session:
            WorkspacePrivacyPolicyService().put(session, actor_user_id=self.fx.user.id,
                policy=WorkspacePrivacyPolicy(workspace_id=self.fx.workspace.id, purpose='Business qualification',
                    legal_basis='legitimate interest', privacy_contact='privacy@example.test'))
            lead = session.scalar(select(LeadModel).where(LeadModel.job_id == job_id))
            contact = ContactPoint.from_email('mario.rossi@official.example.test', source_url='https://official.example.test/',
                collected_at=datetime.now(timezone.utc), extraction_method=ContactExtractionMethod.MAILTO,
                evidence_sha256='a' * 64)
            self.worker.runtime.contacts.persist(session, workspace_id=self.fx.workspace.id, lead_id=lead.id, contacts=[contact])
        return job_id

    def _disable_privacy_policy(self):
        from src.infrastructure.models import WorkspacePrivacyPolicyModel
        with self.fx.database.session(self.fx.workspace.id) as session:
            session.get(WorkspacePrivacyPolicyModel, self.fx.workspace.id).enabled = False

    def test_unrelated_named_contact_policy_failure_does_not_block_generic_job_delivery(self):
        self._named_contact_job()
        response = self.create(kind='audit.website', parameters={'urls': ['https://official.example.test/'], 'max_pages': 1})
        job_id = UUID(response.json()['id']); self.worker.runtime.execute(job_id)
        self._disable_privacy_policy()
        url = self.url + '/' + str(job_id)
        with TestClient(self.client.app, base_url='https://testserver', raise_server_exceptions=False) as client:
            result = client.get(url + '/results', headers=self.headers)
            self.assertEqual(result.status_code, 200, result.text)
            self.assertEqual(result.json()['leads'][0]['contacts'][0]['value'], 'sales@official.example.test')
            export = client.get(url + '/export', headers=self.headers)
            self.assertEqual(export.status_code, 200, export.text)
            rows = repr(list(load_workbook(BytesIO(export.content)).active.values))
            self.assertIn('sales@official.example.test', rows)
            self.assertNotIn('mario.rossi', rows)

    def test_requested_named_contact_policy_failure_returns_safe_denial(self):
        job_id = self._named_contact_job()
        self._disable_privacy_policy()
        url = self.url + '/' + str(job_id)
        with TestClient(self.client.app, base_url='https://testserver', raise_server_exceptions=False) as client:
            for endpoint in ['/results', '/export']:
                with self.subTest(endpoint=endpoint):
                    response = client.get(url + endpoint, headers=self.headers)
                    self.assertEqual(response.status_code, 403, response.text)
                    self.assertEqual(response.json(), {'detail': 'job_contact_policy_denied'})

    def test_named_contact_policy_revocation_between_stream_chunks_stops_delivery(self):
        from src.web.dependencies import commercial_results
        job_id = self._named_contact_job()
        service = commercial_results(self.app_runtime)
        _, _, permission, snapshot = service.export(job_id=job_id, workspace_id=self.fx.workspace.id, auth=self.fx.auth)
        stream = service.stream(b'x' * 150000, job_id=job_id, workspace_id=self.fx.workspace.id,
            auth=self.fx.auth, permission=permission, snapshot=snapshot)
        self.assertEqual(len(next(stream)), 65536)
        self._disable_privacy_policy()
        self.assertEqual(list(stream), [])

    def test_admission_ignores_client_cost_and_persists_real_session_and_quote(self):
        response = self.create(key='same'); self.assertEqual(response.status_code, 202, response.text)
        job_id = UUID(response.json()['id'])
        self.assertEqual(self.publisher.messages, [job_id])
        with self.fx.database.session(self.fx.workspace.id) as session:
            job = session.get(JobModel, job_id)
            self.assertEqual(job.estimated_cost, Decimal('1.8'))
            self.assertEqual(job.auth_session_id, self.fx.auth.session_id)
            self.assertEqual(session.scalar(select(UsageBudgetModel)).reserved, Decimal('1.8'))
        repeat = self.create(key='same'); self.assertEqual(repeat.json()['id'], str(job_id))
        changed = self.create(key='same', parameters={'lat': 1, 'lng': 0, 'keywords': ['dentista']})
        self.assertEqual(changed.status_code, 409)

    def test_unknown_forged_input_and_missing_prices_never_queue(self):
        for kind, payload in [('audit', {}), ('audit.website', {'urls': ['https://site.test/'], 'mfa_verified': True}),
                              ('discover.no_website', {'places': [{'displayName': 'PRIVATE-SENTINEL'}]})]:
            response = self.create(kind=kind, parameters=payload)
            self.assertEqual(response.status_code, 422, response.text)
            self.assertNotIn('PRIVATE-SENTINEL', response.text)
        self.client.app.state.runtime = replace(self.app_runtime, job_prices=None)
        self.assertEqual(self.create().status_code, 503)
        self.assertEqual(self.publisher.messages, [])

    def test_reference_results_and_licensed_export_are_reopened_and_expiry_denies_delivery(self):
        self.fx.activate(('export.no_website',))
        response = self.create(); self.assertEqual(response.status_code, 202, response.text)
        job_id = UUID(response.json()['id']); self.worker.runtime.execute(job_id)
        url = self.url + '/' + str(job_id)
        result = self.client.get(url + '/results', headers=self.headers)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()['references'][0]['place_id'], 'place-1')
        exported = self.client.get(url + '/export', headers=self.headers)
        self.assertEqual(exported.status_code, 200, exported.text)
        book = load_workbook(BytesIO(exported.content))
        self.assertEqual(book.active.cell(2, 1).value, 'place-1')
        self.assertNotIn('PRIVATE GOOGLE NAME', repr(list(book.active.values)))
        self.fx.clock.set(100)
        self.assertEqual(self.client.get(url + '/export', headers=self.headers).status_code, 403)

    def test_revocation_during_export_rendering_returns_no_workbook(self):
        scope, claims = self.fx.activate(('export.no_website',))
        response = self.create(); job_id = UUID(response.json()['id']); self.worker.runtime.execute(job_id)
        from src.application.reference_exports import _serialize_reference_workbook
        def revoke(references):
            content = _serialize_reference_workbook(references)
            with self.fx.database.session(self.fx.workspace.id) as session:
                self.fx.access(session).licenses.revoke_license(scope, claims.license_id)
            return content
        with patch('src.application.reference_exports._serialize_reference_workbook', side_effect=revoke):
            exported = self.client.get(self.url + '/' + str(job_id) + '/export', headers=self.headers)
        self.assertEqual(exported.status_code, 403, exported.text)
        self.assertFalse(exported.content.startswith(b'PK'))

    def test_website_delivery_has_evidence_and_current_contact_suppression(self):
        response = self.create(kind='audit.website', parameters={'urls': ['https://official.example.test/'], 'max_pages': 1})
        self.assertEqual(response.status_code, 202, response.text)
        job_id = UUID(response.json()['id']); self.worker.runtime.execute(job_id)
        url = self.url + '/' + str(job_id) + '/results'
        result = self.client.get(url, headers=self.headers)
        self.assertEqual(result.status_code, 200, result.text)
        lead = result.json()['leads'][0]
        self.assertEqual(lead['attributes']['business_name'], 'Official Clinic')
        self.assertTrue(lead['evidence'][0]['content_sha256'])
        self.assertEqual(len(lead['contacts']), 1)
        with self.fx.database.session(self.fx.workspace.id) as session:
            self.worker.runtime.contacts.suppression.suppress(session, workspace_id=self.fx.workspace.id,
                kind='email', value='sales@official.example.test', scope='workspace', reason='opposition')
        self.assertEqual(self.client.get(url, headers=self.headers).json()['leads'][0]['contacts'], [])

    def test_diagnostic_zip_and_workspace_isolation(self):
        self.fx.activate()
        response = self.create(kind='diagnostic.website', parameters={'url': 'https://official.example.test/', 'max_pages': 1})
        self.assertEqual(response.status_code, 202, response.text)
        job_id = UUID(response.json()['id']); self.worker.runtime.execute(job_id)
        exported = self.client.get(self.url + '/' + str(job_id) + '/export', headers=self.headers)
        self.assertEqual(exported.status_code, 200, exported.text)
        with ZipFile(BytesIO(exported.content)) as archive:
            self.assertEqual(json.loads(archive.read('manifest.json'))['classification'], 'PRIVATE_REDACTED_DIAGNOSTICS')
        other = f'/workspaces/{uuid4()}/jobs/{job_id}/results'
        self.assertIn(self.client.get(other, headers=self.headers).status_code, [403, 404])
        self.fx.clock.set(100)
        self.assertEqual(self.client.get(self.url + '/' + str(job_id) + '/export', headers=self.headers).status_code, 403)

    def test_cancel_and_resume_use_current_session_and_preserve_checkpoint(self):
        response = self.create(); self.assertEqual(response.status_code, 202, response.text)
        job_id = UUID(response.json()['id']); url = self.url + '/' + str(job_id)
        self.assertEqual(self.client.post(url + '/cancel', headers=self.headers).status_code, 200)
        self.assertEqual(self.client.post(url + '/resume', headers=self.headers).status_code, 202)
        self.worker.runtime.execute(job_id)
        self.assertEqual(self.client.get(url, headers=self.headers).json()['state'], 'completed')

    def test_resume_budget_denial_is_controlled_and_preserves_cancelled_state(self):
        response = self.create(); job_id = UUID(response.json()['id']); url = self.url + '/' + str(job_id)
        self.assertEqual(self.client.post(url + '/cancel', headers=self.headers).status_code, 200)
        messages = list(self.publisher.messages)
        with self.fx.database.session(self.fx.workspace.id) as session:
            session.scalar(select(UsageBudgetModel)).hard_limit = Decimal('0.1')
        with TestClient(self.client.app, base_url='https://testserver', raise_server_exceptions=False) as client:
            response = client.post(url + '/resume', headers=self.headers)
            self.assertEqual(response.status_code, 409, response.text)
            self.assertEqual(response.json(), {'detail': 'budget_exceeded'})
            admission = self.create()
            self.assertEqual(admission.status_code, 409)
            self.assertEqual(admission.json(), {'detail': 'budget_exceeded'})
        with self.fx.database.session(self.fx.workspace.id) as session:
            self.assertEqual(session.get(JobModel, job_id).state, 'cancelled')
            self.assertEqual(session.scalar(select(UsageBudgetModel)).reserved, Decimal('0'))
        self.assertEqual(self.publisher.messages, messages)

    def test_reference_stream_rechecks_export_grant_between_chunks(self):
        from src.web.dependencies import commercial_results
        from src.domain.identity import Permission
        scope, claims = self.fx.activate(('export.no_website',))
        response = self.create(); job_id = UUID(response.json()['id']); self.worker.runtime.execute(job_id)
        stream = commercial_results(self.app_runtime).stream(b'x' * 150000, job_id=job_id,
            workspace_id=self.fx.workspace.id, auth=self.fx.auth, permission=Permission.EXPORT_RESULTS)
        self.assertEqual(len(next(stream)), 65536)
        with self.fx.database.session(self.fx.workspace.id) as session:
            self.fx.access(session).licenses.revoke_license(scope, claims.license_id)
        self.assertEqual(len(list(stream)), 0)
        fresh_stream = commercial_results(self.app_runtime).stream(b'x' * 100000, job_id=job_id,
            workspace_id=self.fx.workspace.id, auth=self.fx.auth, permission=Permission.EXPORT_RESULTS)
        self.assertEqual(len(list(fresh_stream)), 0)

    def test_contact_opposition_during_workbook_rendering_denies_stale_export(self):
        response = self.create(kind='audit.website', parameters={'urls': ['https://official.example.test/'], 'max_pages': 1})
        job_id = UUID(response.json()['id']); self.worker.runtime.execute(job_id)
        from src.application.job_results import _verified_workbook
        def oppose(leads):
            content = _verified_workbook(leads)
            with self.fx.database.session(self.fx.workspace.id) as session:
                self.worker.runtime.contacts.suppression.suppress(session, workspace_id=self.fx.workspace.id,
                    kind='email', value='sales@official.example.test', scope='workspace', reason='opposition')
            return content
        with patch('src.application.job_results._verified_workbook', side_effect=oppose):
            result = self.client.get(self.url + '/' + str(job_id) + '/export', headers=self.headers)
        self.assertEqual(result.status_code, 409)
        self.assertFalse(result.content.startswith(b'PK'))

    def test_diagnostic_stream_stops_at_artifact_deadline_before_license_expiry(self):
        from src.domain.feature_licenses import LicenseScope, SubjectKind
        from src.domain.identity import Permission
        from src.web.dependencies import commercial_results
        from tests.license_helpers import license_claims, signed_test_license
        scope = LicenseScope(self.fx.installation, SubjectKind.WORKSPACE_USER, self.fx.user.id, self.fx.workspace.id)
        claims = license_claims(scope=scope, expires_at=10000)
        with self.fx.database.session(self.fx.workspace.id) as session:
            self.fx.access(session).licenses.import_license(scope, signed_test_license(claims, self.fx.private))
        response = self.create(kind='diagnostic.website', parameters={'url': 'https://official.example.test/', 'max_pages': 1, 'retention_hours': 1})
        job_id = UUID(response.json()['id']); self.worker.runtime.execute(job_id)
        content = self.client.get(self.url + '/' + str(job_id) + '/export', headers=self.headers).content
        with ZipFile(BytesIO(content)) as archive: manifest = archive.read('manifest.json')
        output = BytesIO()
        with ZipFile(output, 'w') as archive:
            archive.writestr('manifest.json', manifest); archive.writestr('padding.txt', b'x' * 150000)
        stream = commercial_results(self.app_runtime).stream(output.getvalue(), job_id=job_id,
            workspace_id=self.fx.workspace.id, auth=self.fx.auth, permission=Permission.VIEW_RESULTS)
        self.assertEqual(len(next(stream)), 65536)
        self.fx.clock.set(3610)
        self.assertEqual(len(list(stream)), 0)
