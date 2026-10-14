import copy
import json
import unittest
from dataclasses import replace
from io import BytesIO
from unittest.mock import patch
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import select

from src.application.discovery import DiscoveryPersistenceService
from src.application.provenance import build_verified_lead
from src.domain.audit import AuditValidationError
from src.domain.dom_checks import metric
from src.domain.markup_score import build_score_report, empty_score_report
from src.exporter import DataExporter
from src.infrastructure.models import LeadAttributeModel, LeadModel
from tests.test_audit_evidence_runtime import LEGACY
from tests.test_dom_checks_delivery import workbook_row
from tests.test_dom_checks_runtime import dom_crawl


class MarkupLocalDeliveryTests(unittest.TestCase):
    def test_local_workbook_appends_numeric_index_and_bound_report(self):
        lead = build_verified_lead(dom_crawl(), search_keyword='Studio', audit=LEGACY)
        content = DataExporter.export_bytes([lead], suppression_checker=lambda _: False)
        row = workbook_row(content)
        self.assertEqual(list(row)[-4:], ['Audit Report', 'DOM Report', 'Markup Score', 'Score Report'])
        self.assertEqual(row['Website Score'], 7)
        self.assertEqual(row['Markup Score'], 100)
        self.assertEqual(json.loads(row['Score Report']), lead.score_report)
        sheet = load_workbook(BytesIO(content)).active
        self.assertEqual(sheet.cell(2, len(row) - 1).data_type, 'n')
        self.assertNotIn('PRIVATE_', row['Score Report'])

    def test_zero_is_numeric_and_unavailable_historical_index_is_blank(self):
        lead = build_verified_lead(dom_crawl(1, html='<html><head></head><body></body></html>'),
            search_keyword='Studio', audit=LEGACY)
        content = DataExporter.export_bytes([lead], suppression_checker=lambda _: False)
        row = workbook_row(content)
        self.assertEqual(row['Markup Score'], 0)
        sheet = load_workbook(BytesIO(content)).active
        cell = sheet.cell(2, len(row) - 1)
        self.assertEqual(cell.data_type, 'n'); self.assertNotEqual(cell.fill.fgColor.rgb, '00FEE2E2')
        historical = replace(lead, score_report=None)
        row = workbook_row(DataExporter.export_bytes([historical], suppression_checker=lambda _: False))
        self.assertIsNone(row['Markup Score'])
        self.assertEqual(json.loads(row['Score Report']), empty_score_report())

    def test_local_export_refuses_score_that_does_not_match_current_dom(self):
        lead = build_verified_lead(dom_crawl(), search_keyword='Studio', audit=LEGACY)
        bad = copy.deepcopy(lead.score_report); bad['score'] = 0
        with self.assertRaisesRegex(AuditValidationError, '^score_report_invalid$'):
            DataExporter.export_bytes([replace(lead, score_report=bad)], suppression_checker=lambda _: False)
        with self.assertRaisesRegex(AuditValidationError, '^score_report_invalid$'):
            DataExporter.export_bytes([replace(lead, dom_report=None)], suppression_checker=lambda _: False)


class MarkupServerDeliveryTests(unittest.TestCase):
    def setUp(self):
        from tests.test_dom_checks_delivery import DomServerDeliveryTests
        self.dom_fx = DomServerDeliveryTests(); self.dom_fx.setUp(); self.addCleanup(self.dom_fx.doCleanups)
        self.fx = self.dom_fx.fx

    def attribute(self, session, field):
        return session.scalar(select(LeadAttributeModel).where(LeadAttributeModel.field_name == field))

    def test_sdk_worker_sql_http_and_server_xlsx_share_separate_numeric_index(self):
        url = self.dom_fx.dom_job()
        response = self.fx.client.get(url + '/results', headers=self.fx.headers)
        self.assertEqual(response.status_code, 200, response.text)
        lead = response.json()['leads'][0]; report = lead['score_report']
        self.assertEqual(report, build_score_report(lead['dom_report']))
        self.assertEqual(report['score'], 100.0)
        self.assertNotIn('score_report', lead['attributes'])
        self.assertEqual(lead['attributes']['website_score'], '7')
        with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
            attribute = self.attribute(session, 'score_report')
            self.assertEqual(json.loads(attribute.value_text), report)
            self.assertEqual(attribute.data_classification, 'website_markup_score')
        export = self.fx.client.get(url + '/export', headers=self.fx.headers)
        self.assertEqual(export.status_code, 200, export.text)
        row = workbook_row(export.content)
        self.assertEqual(list(row)[-4:], ['audit_report', 'dom_report', 'markup_score', 'score_report'])
        self.assertEqual(row['markup_score'], 100)
        self.assertEqual(json.loads(row['score_report']), report)
        sheet = load_workbook(BytesIO(export.content)).active
        self.assertEqual(sheet.cell(2, len(row) - 1).data_type, 'n')
        self.assertNotIn('PRIVATE_', row['score_report'])

    def test_corrupt_json_score_bool_or_dom_metadata_fails_without_sql_echo(self):
        url = self.dom_fx.dom_job()
        with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
            original = self.attribute(session, 'score_report').value_text
        for changes in ('{PRIVATE_SQL_SENTINEL', 'wrong_score', 'bool_score'):
            with self.subTest(changes=changes):
                value = changes
                if changes != '{PRIVATE_SQL_SENTINEL':
                    bad = json.loads(original); bad['score'] = True if changes == 'bool_score' else 0
                    value = json.dumps(bad)
                with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
                    self.attribute(session, 'score_report').value_text = value
                response = self.fx.client.get(url + '/results', headers=self.fx.headers)
                self.assertEqual(response.json(), {'detail': 'job_score_report_invalid'})
                self.assertNotIn('PRIVATE_SQL_SENTINEL', response.text)
        with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
            self.attribute(session, 'score_report').value_text = original
            dom_attribute = self.attribute(session, 'dom_report'); dom = json.loads(dom_attribute.value_text)
            dom['pages'][0]['retrieved_at'] = '2000-01-01T00:00:00+00:00'
            dom_attribute.value_text = json.dumps(dom)
        response = self.fx.client.get(url + '/export', headers=self.fx.headers)
        self.assertEqual(response.json(), {'detail': 'job_score_report_invalid'})
        with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
            session.delete(self.attribute(session, 'dom_report'))
        response = self.fx.client.get(url + '/results', headers=self.fx.headers)
        self.assertEqual(response.json(), {'detail': 'job_score_report_invalid'})

    def test_historical_missing_score_remains_unavailable_even_with_dom_present(self):
        url = self.dom_fx.dom_job()
        with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
            session.delete(self.attribute(session, 'score_report'))
        response = self.fx.client.get(url + '/results', headers=self.fx.headers)
        self.assertEqual(response.status_code, 200, response.text)
        lead = response.json()['leads'][0]
        self.assertEqual(lead['score_report'], empty_score_report())
        self.assertEqual(lead['dom_report']['status'], 'complete')
        self.assertTrue(all(isinstance(value, str) for value in lead['attributes'].values()))
        row = workbook_row(self.fx.client.get(url + '/export', headers=self.fx.headers).content)
        self.assertIsNone(row['markup_score'])

    def test_valid_score_change_during_workbook_render_is_blocked_by_snapshot_guard(self):
        from src.application import job_results
        url = self.dom_fx.dom_job(); original_renderer = job_results._verified_workbook
        def changed(leads):
            content = original_renderer(leads)
            with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
                dom_attribute = self.attribute(session, 'dom_report'); dom = json.loads(dom_attribute.value_text)
                dom['pages'][0]['checks']['title_text'] = metric(0, 1)
                dom_attribute.value_text = json.dumps(dom)
                self.attribute(session, 'score_report').value_text = json.dumps(build_score_report(dom))
            return content
        with patch('src.application.job_results._verified_workbook', side_effect=changed):
            response = self.fx.client.get(url + '/export', headers=self.fx.headers)
        self.assertEqual(response.json(), {'detail': 'job_result_changed'})

    def test_persistence_rejects_forged_score_before_delivering_a_new_lead(self):
        url = self.dom_fx.dom_job(); job_id = UUID(url.rsplit('/', 1)[1])
        lead = build_verified_lead(dom_crawl(), search_keyword='Studio', audit=LEGACY)
        bad = copy.deepcopy(lead.score_report); bad['score'] = -1
        with self.assertRaisesRegex(AuditValidationError, '^score_report_invalid$'):
            with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
                DiscoveryPersistenceService().persist_verified_lead(session, workspace_id=self.fx.fx.workspace.id,
                    job_id=job_id, lead=replace(lead, score_report=bad))
        with self.fx.fx.database.session(self.fx.fx.workspace.id) as session:
            self.assertEqual(len(session.scalars(select(LeadModel).where(LeadModel.job_id == job_id)).all()), 1)
