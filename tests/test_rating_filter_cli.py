import contextlib
import unittest
from io import StringIO
from unittest.mock import patch

from openpyxl import load_workbook

from main import main, LeadHunterOrchestrator, create_orchestrator
from src.application.container import ApplicationContainer
from src.domain.rating_filters import RatingFilterCriteria
from src.scraper import LeadScraper
from tests.rating_filter_helpers import RatingFilterFixture
from tests.test_rating_filter_discovery import Transport, make_provider, place
from tests.test_compliant_pipeline import _Auditor, _Crawler


class RatingFilterCliTests(unittest.TestCase):
    def setUp(self):
        self.fx = RatingFilterFixture(self, features=('discovery.rating_filters', 'export.no_website'))
        self.output = StringIO()
        redirect = contextlib.redirect_stdout(self.output); redirect.__enter__()
        self.addCleanup(redirect.__exit__, None, None, None)
        self.destination = self.fx.root / 'report.xlsx'
        self.args = ['--lat', '45', '--lng', '9', '--keywords', 'dentista', '--out', str(self.destination)]
        env = patch.dict('os.environ', {'GOOGLE_API_KEY': 'test-only', 'OPENROUTER_API_KEY': 'test-only'})
        env.start(); self.addCleanup(env.stop)
        self.http = Transport()
        self.factory = patch('src.application.container.LeadScraper',
            side_effect=lambda **kwargs: LeadScraper(provider=make_provider(self.http)))
        self.scraper_factory = self.factory.start(); self.addCleanup(self.factory.stop)

    def test_explicit_flag_applies_defaults_in_both_modes(self):
        self.fx.activate()
        for mode in ['no_website', 'with_website']:
            with self.subTest(mode=mode):
                row = place('qualified', websiteUri='https://official.example.test') if mode == 'with_website' else place('qualified')
                self.http.places = [row]
                crawler = _Crawler()
                original = crawler.crawl
                async def no_contacts(url):
                    result = await original(url); result.contacts = []; return result
                crawler.crawl = no_contacts
                with patch.object(ApplicationContainer, 'build_auditor', return_value=_Auditor()), patch('main.HybridCrawler', return_value=crawler), patch('main.filter_by_business_age', return_value=True), patch('main.filter_franchise', return_value=False), patch('main.create_orchestrator', wraps=create_orchestrator) as constructor:
                    self.assertEqual(main(self.args + ['--mode', mode, '--rating-filters']), 0)
                self.assertEqual(constructor.call_args.kwargs['rating_criteria'], RatingFilterCriteria(3.9, 100))
                self.assertEqual(set(self.http.calls[-1]['headers']['X-Goog-FieldMask'].split(',')) & {'places.rating', 'places.userRatingCount'},
                                 {'places.rating', 'places.userRatingCount'})

    def test_thresholds_without_flag_or_special_modes_fail_before_any_io(self):
        options = [['--min-rating', '3.9'], ['--max-reviews', '100']]
        for special in [['--test-url', 'https://official.test'], ['--gui'], ['--examples']]:
            options.append(['--rating-filters'] + special)
        with patch('main.ApplicationSettings.from_environment') as settings:
            for option in options:
                with self.subTest(option=option), contextlib.redirect_stderr(StringIO()), self.assertRaises(SystemExit) as caught:
                    main(self.args + option)
                self.assertEqual(caught.exception.code, 2)
            settings.assert_not_called()
        self.scraper_factory.assert_not_called()

    def test_invalid_criteria_and_license_prevent_all_side_effects(self):
        for option in [['--min-rating', 'nan'], ['--min-rating', 'inf'], ['--min-rating', '-1'],
                       ['--max-reviews', '0'], ['--max-reviews', '2147483648']]:
            with self.subTest(option=option):
                self.assertEqual(main(self.args + ['--rating-filters'] + option), 2)
        self.assertEqual(main(self.args + ['--rating-filters']), 2)
        self.assertEqual(len(self.http.calls), 0)
        self.scraper_factory.assert_not_called()
        self.assertFalse(self.destination.exists())

    def test_reference_delivery_denied_preserves_destination(self):
        self.fx.activate(); self.destination.write_bytes(b'previous-export')
        from src.domain.place_references import project_google_place_references
        def revoke(results):
            references = project_google_place_references(results)
            self.fx.licenses.revoke_license(self.fx.scope, self.fx.claims.license_id)
            return references
        with patch('main.project_google_place_references', side_effect=revoke):
            self.assertEqual(main(self.args + ['--rating-filters', '--export-references']), 2)
        self.assertEqual(self.destination.read_bytes(), b'previous-export')
        self.assertNotIn('qualified', self.output.getvalue())

    def test_valid_filtered_reference_export_has_only_two_columns(self):
        self.fx.activate()
        self.assertEqual(main(self.args + ['--rating-filters', '--export-references']), 0)
        self.assertEqual(load_workbook(self.destination).active.max_column, 2)

    def test_delivery_denied_or_interrupted_has_no_partial_export(self):
        self.fx.activate(); self.destination.write_bytes(b'previous-export')
        for error in [KeyboardInterrupt(), RuntimeError('provider-secret-sentinel')]:
            with self.subTest(error=type(error).__name__), patch.object(ApplicationContainer, 'build_auditor', return_value=_Auditor()), patch.object(LeadHunterOrchestrator, 'run', side_effect=error):
                self.assertEqual(main(self.args + ['--mode', 'with_website', '--rating-filters']), 1)
            self.assertEqual(self.destination.read_bytes(), b'previous-export')
            self.assertNotIn('provider-secret-sentinel', self.output.getvalue())

    def test_dependency_construction_failure_is_redacted(self):
        self.fx.activate()
        with patch.object(ApplicationContainer, 'build_auditor', side_effect=RuntimeError('provider-secret-sentinel')):
            self.assertEqual(main(self.args + ['--mode', 'with_website', '--rating-filters']), 1)
        self.assertNotIn('provider-secret-sentinel', self.output.getvalue())
        self.assertFalse(self.destination.exists())

    def test_base_needs_no_license_configuration(self):
        with patch.object(ApplicationContainer, 'build_local_license_service', side_effect=AssertionError('base licensing I/O')):
            self.assertEqual(main(self.args), 0)
        self.assertEqual(self.http.calls[0]['headers']['X-Goog-FieldMask'].count('places.rating'), 0)
