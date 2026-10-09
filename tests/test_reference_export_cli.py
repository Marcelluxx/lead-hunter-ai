import contextlib
from datetime import datetime, timezone
import io
from pathlib import Path
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from main import main
from src.domain.discovery import ProviderAttribution, TransientCandidate
from tests.reference_export_helpers import ReferenceExportFixture


class ReferenceExportCliTests(unittest.TestCase):
    def setUp(self):
        self.fx = ReferenceExportFixture(self)
        self.output = io.StringIO()
        redirect = contextlib.redirect_stdout(self.output); redirect.__enter__()
        self.addCleanup(redirect.__exit__, None, None, None)
        self.output_dir = self.fx.root / 'output'
        self.destination = self.output_dir / 'references.xlsx'
        output_patch = patch('main.OUTPUT_DIR', str(self.output_dir))
        output_patch.start(); self.addCleanup(output_patch.stop)
        self.provider_calls = 0
        self.attribution = ProviderAttribution('google_places', 'Google', 'https://example.test/terms',
                                              'https://example.test/privacy')
        self.results = [TransientCandidate('google_places', '0007', 'discard-business', None,
                                           datetime.now(timezone.utc), self.attribution)]
        testcase = self
        class Scraper:
            attribution = testcase.attribution
            def get_city_name(self, lat, lng):
                testcase.provider_calls += 1
                return 'UserArea'
        class Orchestrator:
            scraper = Scraper()
            def run(self, lat, lng, keywords, **kwargs):
                testcase.provider_calls += 1
                return testcase.results
        constructor = patch('main.create_orchestrator', return_value=Orchestrator())
        constructor.start(); self.addCleanup(constructor.stop)
        self.search_args = ['--lat', '45.4642', '--lng', '9.1900', '--keywords', 'ristorante',
                            '--mode', 'no_website']

    def test_parser_rejects_reference_flag_with_website(self):
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as caught:
            main(['--mode', 'with_website', '--export-references'])
        self.assertEqual(caught.exception.code, 2)
        self.assertEqual(self.provider_calls, 0)

    def test_denied_preflight_makes_zero_provider_calls(self):
        self.assertEqual(main(self.search_args + ['--export-references']), 2)
        self.assertEqual(self.provider_calls, 0)
        self.assertFalse(self.output_dir.exists())
        self.assertIn('license_missing', self.output.getvalue())

    def test_valid_flag_saves_two_column_workbook(self):
        self.fx.activate()
        self.assertEqual(main(self.search_args + ['--export-references', '--out', str(self.destination)]), 0)
        sheet = load_workbook(self.destination).active
        self.assertEqual(sheet.max_column, 2)
        self.assertEqual(sheet['A2'].value, '0007')
        self.assertIn('1 riferimenti esportati', self.output.getvalue())
        self.assertNotIn(self.fx.token, self.output.getvalue())

    def test_failed_write_has_no_success_message(self):
        self.fx.activate()
        self.output_dir.mkdir(); self.destination.write_bytes(b'previous')
        with patch('src.application.reference_exports.os.replace', side_effect=PermissionError('discard-secret')):
            self.assertEqual(main(self.search_args + ['--export-references', '--out', str(self.destination)]), 2)
        self.assertNotIn('esportati', self.output.getvalue())
        self.assertIn('reference_export_write_failed', self.output.getvalue())
        self.assertNotIn('discard-secret', self.output.getvalue())
        self.assertEqual(self.destination.read_bytes(), b'previous')
        self.assertEqual(list(self.output_dir.iterdir()), [self.destination])

    def test_base_search_needs_no_license_and_writes_no_workbook(self):
        self.assertEqual(main(self.search_args), 0)
        self.assertGreater(self.provider_calls, 0)
        self.assertFalse(self.output_dir.exists())
        self.assertIn('risultati transitori', self.output.getvalue())

    def test_empty_search_creates_no_file(self):
        self.fx.activate(); self.results = []
        self.assertEqual(main(self.search_args + ['--export-references', '--out', str(self.destination)]), 0)
        self.assertFalse(self.destination.exists())
        self.assertNotIn('esportati', self.output.getvalue())

    def test_non_google_results_are_not_exported(self):
        from dataclasses import replace
        self.fx.activate(); self.results = [replace(self.results[0], provider='other')]
        self.assertEqual(main(self.search_args + ['--export-references', '--out', str(self.destination)]), 0)
        self.assertFalse(self.destination.exists())

    def test_bare_filename_and_default_name_use_output_directory(self):
        self.fx.activate()
        self.assertEqual(main(self.search_args + ['--export-references', '--out', 'bare.xlsx']), 0)
        self.assertEqual(load_workbook(self.output_dir / 'bare.xlsx').active['A2'].value, '0007')
        self.assertEqual(main(self.search_args + ['--export-references']), 0)
        files = list(self.output_dir.glob('Lead_Hunter_UserArea_*.xlsx'))
        self.assertEqual(len(files), 1)
        self.assertEqual(load_workbook(files[0]).active['A2'].value, '0007')
