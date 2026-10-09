import json
import os
import unittest
from dataclasses import replace
from io import BytesIO
from itertools import repeat
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile

from openpyxl import load_workbook
from sqlalchemy import delete, update

from src.application.feature_access import FeatureAccessService, server_feature_context
from src.application.feature_licenses import LicenseService
from src.application.reference_exports import ReferenceExportService, _serialize_reference_workbook
from src.domain.feature_licenses import LicenseError
from src.domain.place_references import GooglePlaceReference, ReferenceExportError, project_google_place_references
from src.infrastructure.license_repository import SqlLicenseRepository
from src.infrastructure.models import UserModel, WorkspaceModel, WorkspaceMembershipModel
from tests.license_helpers import license_claims, seed_license_subject, signed_test_license
from tests.platform_helpers import platform_fixture
from tests.reference_export_helpers import ReferenceExportFixture, AvailableExportCatalog


class ReferenceExportTests(unittest.TestCase):
    def setUp(self):
        self.fx = ReferenceExportFixture(self)
        self.fx.activate()
        self.service = self.fx.service()
        self.refs = tuple(map(GooglePlaceReference, ['0007', 'Ab', 'ab', '0007', 'A' * 500]))
        self.output_dir = self.fx.root / 'output'
        self.output_dir.mkdir()
        self.destination = self.output_dir / 'references.xlsx'

    def test_real_workbook_has_only_reference_content(self):
        data = self.service.export_bytes(self.refs)
        book = load_workbook(BytesIO(data)); sheet = book['Riferimenti']
        self.assertEqual(book.sheetnames, ['Riferimenti'])
        self.assertEqual(tuple(c.value for c in sheet[1]), ('Place ID', 'Link Google Maps'))
        self.assertEqual([sheet.cell(i, 1).value for i in range(2, 6)], ['0007', 'Ab', 'ab', 'A' * 500])
        self.assertEqual(sheet.freeze_panes, 'A2'); self.assertEqual(sheet.auto_filter.ref, 'A1:B5')
        self.assertEqual(sheet['B2'].value,
                         'https://www.google.com/maps/search/?api=1&query=attivit%C3%A0&query_place_id=0007')
        self.assertTrue(all(c.data_type == 's' and c.comment is None for row in sheet for c in row))
        self.assertTrue(all(sheet.cell(i, 2).hyperlink.target == sheet.cell(i, 2).value for i in range(2, 6)))
        self.assertEqual(sheet.sheet_state, 'visible')
        with ZipFile(BytesIO(data)) as archive:
            for name in archive.namelist():
                member = archive.read(name)
                for forbidden in [b'discard-business', b'discard-keyword', b'discard-phone', self.fx.token.encode()]:
                    self.assertNotIn(forbidden, member)
            self.assertFalse(any('comments' in name or 'externalLinks' in name for name in archive.namelist()))

    def test_projection_content_never_reaches_workbook_zip(self):
        from datetime import datetime, timezone
        from src.domain.discovery import TransientCandidate, ProviderAttribution
        candidate = TransientCandidate('google_places', 'ChIJ_1', 'discard-business', None,
            datetime.now(timezone.utc), ProviderAttribution('google_places', 'discard-keyword',
                                                           'https://discard-phone', 'https://discard-prompt'))
        data = self.service.export_bytes(project_google_place_references([candidate]))
        with ZipFile(BytesIO(data)) as archive:
            contents = b''.join(archive.read(name) for name in archive.namelist())
        for marker in [b'discard-business', b'discard-keyword', b'discard-phone', b'discard-prompt']:
            self.assertNotIn(marker, contents)

    def test_context_factory_is_called_for_generation_and_delivery(self):
        contexts = []
        original = self.service.context_factory
        def fresh():
            context = original(); contexts.append(context); return context
        service = ReferenceExportService(self.service.access, fresh)
        service.export_bytes(self.refs)
        self.assertEqual(len(contexts), 2)
        self.assertIsNot(contexts[0], contexts[1])

    def test_delivery_rechecks_after_expiry_revoke_or_renewal(self):
        changes = [('expiry', 'license_expired'), ('revoke', 'license_revoked'), ('renew', 'feature_not_granted')]
        for scenario, code in changes:
            with self.subTest(scenario=scenario):
                fx = ReferenceExportFixture(self); fx.activate(); service = fx.service()
                def serialize(refs):
                    data = _serialize_reference_workbook(refs)
                    if scenario == 'expiry':
                        fx.clock.set(100)
                    elif scenario == 'revoke':
                        fx.licenses.revoke_license(fx.scope, fx.claims.license_id)
                    else:
                        fx.activate(replace(fx.claims, license_id=uuid4(), features=('diagnostics.full',)))
                    return data
                with patch('src.application.reference_exports._serialize_reference_workbook', side_effect=serialize):
                    with self.assertRaises(LicenseError) as caught:
                        service.export_bytes(self.refs)
                self.assertEqual(caught.exception.code, code)

    def test_direct_call_denies_missing_tampered_expired_revoked_other_subject_and_missing_feature(self):
        for scenario, code in [('missing', 'license_missing'), ('tampered', 'license_invalid'),
                               ('expired', 'license_expired'), ('revoked', 'license_revoked'),
                               ('other', 'license_subject_mismatch'), ('feature', 'feature_not_granted')]:
            with self.subTest(scenario=scenario):
                fx = ReferenceExportFixture(self)
                if scenario != 'missing':
                    fx.activate()
                if scenario in ('tampered', 'other'):
                    state_path = fx.root / 'state' / 'state.json'
                    state = json.loads(state_path.read_text())
                    if scenario == 'tampered':
                        state['grant']['token'] = 'tampered.test.token'
                    else:
                        other_identity = uuid4()
                        state['grant']['token'] = signed_test_license(
                            replace(fx.claims, scope=replace(fx.scope, installation_id=other_identity,
                                                            subject_id=other_identity)), fx.private)
                    state_path.write_text(json.dumps(state))
                elif scenario == 'expired':
                    fx.clock.set(100)
                elif scenario == 'revoked':
                    fx.licenses.revoke_license(fx.scope, fx.claims.license_id)
                elif scenario == 'feature':
                    fx.activate(replace(fx.claims, license_id=uuid4(), features=('diagnostics.full',)))
                with self.assertRaises(LicenseError) as caught:
                    fx.service().export_bytes(self.refs)
                self.assertEqual(caught.exception.code, code)

    def test_not_before_and_clock_regression_remain_enforced(self):
        state_path = self.fx.root / 'state' / 'state.json'
        state = json.loads(state_path.read_text())
        state['grant']['token'] = signed_test_license(replace(self.fx.claims, not_before=20), self.fx.private)
        state_path.write_text(json.dumps(state))
        with self.assertRaises(LicenseError) as caught:
            self.service.export_bytes(self.refs)
        self.assertEqual(caught.exception.code, 'license_not_yet_valid')
        self.fx.clock.set(1000)
        self.fx.activate(replace(self.fx.claims, license_id=uuid4(), expires_at=2000, not_before=1000))
        self.fx.clock.set(600)
        with self.assertRaises(LicenseError) as caught:
            self.service.export_bytes(self.refs)
        self.assertEqual(caught.exception.code, 'license_clock_regression')

    def test_empty_invalid_or_over_limit_never_creates_file(self):
        for refs, code in [([], 'reference_export_empty'), ([{'place_id': 'A'}], 'reference_export_invalid'),
                           (repeat(GooglePlaceReference('A'), 10001), 'reference_export_limit')]:
            with self.subTest(code=code), self.assertRaises(ReferenceExportError) as caught:
                self.service.save(refs, self.destination)
            self.assertEqual(caught.exception.code, code)
            self.assertEqual(list(self.output_dir.iterdir()), [])

    def test_successful_save_reopens_workbook_and_returns_unique_count(self):
        self.assertEqual(self.service.save(iter(self.refs), self.destination), 4)
        self.assertEqual(load_workbook(self.destination).active['A2'].value, '0007')
        self.assertEqual(list(self.output_dir.iterdir()), [self.destination])

    def test_atomic_save_preserves_old_file_on_denial_or_replace_error(self):
        for scenario, code in [('expire', 'license_expired'), ('replace', 'reference_export_write_failed')]:
            with self.subTest(scenario=scenario):
                fx = ReferenceExportFixture(self); fx.activate(); service = fx.service()
                self.destination.write_bytes(b'previous')
                original = os.fsync
                def fsync(descriptor):
                    original(descriptor); fx.clock.set(100)
                replacement = patch('src.application.reference_exports.os.replace',
                                    side_effect=PermissionError('discard-secret'))
                timing = patch('src.application.reference_exports.os.fsync', side_effect=fsync)
                with (timing if scenario == 'expire' else replacement):
                    with self.assertRaises((LicenseError, ReferenceExportError)) as caught:
                        service.save(self.refs, self.destination)
                self.assertEqual(str(caught.exception), code)
                self.assertEqual(self.destination.read_bytes(), b'previous')
                self.assertEqual(list(self.output_dir.iterdir()), [self.destination])

    def test_write_failure_is_redacted_and_cleans_temp(self):
        with patch('src.application.reference_exports.os.fsync', side_effect=OSError('discard-secret')):
            with self.assertRaises(ReferenceExportError) as caught:
                self.service.save(self.refs, self.destination)
        self.assertEqual(str(caught.exception), 'reference_export_write_failed')
        self.assertEqual(list(self.output_dir.iterdir()), [])

    def test_interrupt_cleans_temp_and_preserves_old_file(self):
        self.destination.write_bytes(b'previous')
        with patch('src.application.reference_exports.os.fsync', side_effect=KeyboardInterrupt):
            with self.assertRaises(KeyboardInterrupt):
                self.service.save(self.refs, self.destination)
        self.assertEqual(self.destination.read_bytes(), b'previous')
        self.assertEqual(list(self.output_dir.iterdir()), [self.destination])

    def test_server_membership_is_fresh_at_delivery(self):
        for change in ['viewer', 'remove', 'inactive_user', 'inactive_workspace']:
            with self.subTest(change=change):
                db, *_ = platform_fixture(); self.addCleanup(db.engine.dispose)
                scope = seed_license_subject(db)
                claims = license_claims(scope=scope, features=('export.no_website',))
                with db.session(scope.workspace_id) as session:
                    licenses = LicenseService(SqlLicenseRepository(session, scope.installation_id),
                                              self.fx.licenses.verifier, self.fx.clock)
                    licenses.import_license(scope, signed_test_license(claims, self.fx.private))
                with db.session(scope.workspace_id) as session:
                    licenses = LicenseService(SqlLicenseRepository(session, scope.installation_id),
                                              self.fx.licenses.verifier, self.fx.clock)
                    auth = SimpleNamespace(user_id=scope.subject_id, mfa_verified=False)
                    service = ReferenceExportService(FeatureAccessService(licenses, AvailableExportCatalog()),
                        lambda: server_feature_context(session, auth=auth, workspace_id=scope.workspace_id,
                                                       installation_id=scope.installation_id))
                    self.assertEqual(load_workbook(BytesIO(service.export_bytes(self.refs))).active['A2'].value, '0007')
                    def serialize(refs):
                        data = _serialize_reference_workbook(refs)
                        with db.session() as writer:
                            member = (WorkspaceMembershipModel.user_id == scope.subject_id)
                            if change == 'viewer':
                                writer.execute(update(WorkspaceMembershipModel).where(member).values(role='viewer'))
                            elif change == 'remove':
                                writer.execute(delete(WorkspaceMembershipModel).where(member))
                            elif change == 'inactive_user':
                                writer.execute(update(UserModel).where(UserModel.id == scope.subject_id).values(is_active=False))
                            else:
                                writer.execute(update(WorkspaceModel).where(WorkspaceModel.id == scope.workspace_id).values(is_active=False))
                        return data
                    with patch('src.application.reference_exports._serialize_reference_workbook', side_effect=serialize):
                        with self.assertRaises(LicenseError) as caught:
                            service.export_bytes(self.refs)
                    self.assertEqual(caught.exception.code, 'feature_role_denied')
