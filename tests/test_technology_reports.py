import asyncio
import unittest
from hashlib import sha256
from io import BytesIO
from uuid import UUID
from openpyxl import load_workbook

from src.application.provenance import build_verified_lead
from src.domain.crawl import CrawlResult, CrawlStatus
from src.domain.provenance import VerifiedLead
from src.exporter import DataExporter
from tests.test_compliant_pipeline import _Crawler

TECH_HTML = '<html><head><title>Studio Aurora</title><meta name="generator" content="WordPress 6"></head><body>PRIVATE_DOM_SENTINEL</body></html>'


class TechnologyReportTests(unittest.TestCase):
    def crawl(self):
        crawl = asyncio.run(_Crawler().crawl('https://official.example.test/'))
        crawl.raw_html_home = TECH_HTML
        return crawl

    def test_report_uses_observed_technology_and_dom_hash_instead_of_llm_guess(self):
        crawl = self.crawl()
        lead = build_verified_lead(crawl, search_keyword='Studio', audit={'framework': 'Inventato', 'website_score': 7})
        self.assertEqual(lead.framework, 'WordPress')
        self.assertTrue(lead.technology_signals)
        self.assertEqual(lead.technology_rules_version, '1')
        self.assertEqual(lead.provenance['framework'].evidence_sha256, sha256(TECH_HTML.encode()).hexdigest())
        self.assertEqual(lead.provenance['website_score'].evidence_sha256, crawl.valid_evidence[0].content_sha256)
        record = repr(lead.to_export_record())
        self.assertNotIn('PRIVATE_DOM_SENTINEL', record)
        self.assertNotIn('Inventato', record)

    def test_unknown_technology_and_changed_dom_do_not_create_fake_proof(self):
        crawl = self.crawl(); crawl.raw_html_home = '<p>WordPress tutorial</p>'
        lead = build_verified_lead(crawl, search_keyword='Studio', audit={'framework': 'Inventato'})
        self.assertEqual(lead.framework, 'Non rilevato')
        self.assertEqual(lead.technology_signals, ())
        crawl.raw_html_home_sha256 = 'a' * 64
        with self.assertRaisesRegex(ValueError, '^crawl_dom_evidence_changed$'):
            build_verified_lead(crawl, search_keyword='Studio')

    def test_local_workbook_exposes_public_technology_signals_without_private_dom(self):
        lead = build_verified_lead(self.crawl(), search_keyword='Studio', audit={'framework': 'Inventato'})
        book = load_workbook(BytesIO(DataExporter.export_bytes([lead], suppression_checker=lambda _: False)))
        row = dict(zip([x.value for x in book.active[1]], [x.value for x in book.active[2]]))
        self.assertEqual(row['Framework'], 'WordPress')
        self.assertIn('generator_meta', row['Technology Signals'])
        self.assertEqual(row['Technology Rules'], '1')
        self.assertNotIn('PRIVATE_DOM_SENTINEL', repr(row))

    def test_new_crawl_fields_preserve_previous_positional_status_argument(self):
        result = CrawlResult('https://site.test/', '', {}, [], [], '', False, CrawlStatus.SUCCESS)
        self.assertIs(result.status, CrawlStatus.SUCCESS)
        lead = VerifiedLead('Site', 'Studio', 'https://site.test/', (), 7, 'CMS', 'Diagnosis', 'Brief', 'Message', {})
        self.assertEqual(lead.diagnosis, 'Diagnosis')


class TechnologyServerReportTests(unittest.TestCase):
    def test_http_worker_persistence_and_xlsx_reopen_use_same_observed_technology(self):
        from tests.test_commercial_job_api import CommercialJobApiTests
        fx = CommercialJobApiTests(); fx.setUp(); self.addCleanup(fx.doCleanups)
        class Crawler(_Crawler):
            async def crawl(self, url):
                result = await super().crawl(url); result.raw_html_home = TECH_HTML
                return result
        fx.worker.runtime.crawler_factory = lambda **kwargs: Crawler()
        response = fx.create(kind='audit.website', parameters={'urls': ['https://official.example.test/'], 'max_pages': 1})
        self.assertEqual(response.status_code, 202, response.text)
        job_id = UUID(response.json()['id']); fx.worker.runtime.execute(job_id)
        url = fx.url + '/' + str(job_id)
        response = fx.client.get(url + '/results', headers=fx.headers)
        self.assertEqual(response.status_code, 200, response.text)
        lead = response.json()['leads'][0]
        self.assertEqual(lead['attributes']['framework'], 'WordPress')
        self.assertIn('generator_meta', lead['attributes']['technology_signals'])
        self.assertEqual(lead['attributes']['technology_rules_version'], '1')
        evidence_id = lead['provenance']['framework']['evidence_id']
        evidence = next(x for x in lead['evidence'] if x['id'] == evidence_id)
        self.assertEqual(evidence['content_sha256'], sha256(TECH_HTML.encode()).hexdigest())
        response = fx.client.get(url + '/export', headers=fx.headers)
        self.assertEqual(response.status_code, 200, response.text)
        book = load_workbook(BytesIO(response.content))
        row = dict(zip([x.value for x in book.active[1]], [x.value for x in book.active[2]]))
        self.assertEqual(row['framework'], 'WordPress')
        self.assertIn('generator_meta', row['technology_signals'])
        self.assertNotIn('PRIVATE_DOM_SENTINEL', repr(row))
